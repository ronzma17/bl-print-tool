#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
提单打印工具 - 引擎
不要直接双击这个文件,请双击 START_HERE.bat

流程:
  解压 zip -> 按提单号配对 -> 合并成双面打印就绪的 PDF -> 原始文件归档
  下次再跑时, 上一批的成品 PDF 会自动挪进同一个归档文件夹, 不用手动删

每票 = [Transport B/L 全部页] + [Waybill 第1页] = 2 页 = 双面打印 1 张纸
        正面 B/L, 反面 Waybill

文件识别:
    短文件名  CHN1234567                -> Waybill (只取第1页)
    长文件名  CHN1122211 1111A1111111   -> Transport B/L
    同时读 PDF 内容交叉验证

依赖: pip install pypdf
"""

import argparse
import csv
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    from pypdf import PdfReader, PdfWriter
except ImportError:
    sys.exit("缺少 pypdf,请先执行:  pip install pypdf")


# ============ 可调参数 ============
KEY_PATTERN = re.compile(r"[A-Z]{2,4}\d{6,}")   # 提单号格式, 如 CHN3966555
WAYBILL_PAGES = [1]                             # Waybill 只打第几页
BL_HINT = "COMBINED TRANSPORT"
WB_HINT = "WAYBILL"
MARKER = "_batch.txt"                           # 记录本批归档到哪个文件夹
# 跑之前要清掉的上一批产物。报告改过格式(csv -> xlsx), 所以这里用前缀匹配,
# 写死文件名的话改一次格式就漏一个, 旧文件会一直赖在输出文件夹里。
EXTRA_OUT = ("_需人工处理.pdf", "_多出来的PREARRIVAL.pdf")
OUT_PREFIX = ("PRINT", "_报告")
# ==================================


CONTAINER_PATTERN = re.compile(r"\b[A-Z]{4}\d{7}\b")  # 柜号, 别当成提单号
HBL_PATTERN = re.compile(r"\d{4}[A-Z]\d{5,8}")        # HBL, 如 2609A338346 / AIC2609A338346


def writable_path(p: Path) -> Path:
    """Windows 上文件被 Excel / Adobe 开着时无法覆盖, 会抛 PermissionError。
       与其整个脚本崩掉, 不如换个带时间戳的名字存下来, 至少数据不丢。"""
    try:
        with open(p, "ab"):
            pass
        return p
    except PermissionError:
        alt = p.with_name(f"{p.stem}_{datetime.now():%H%M%S}{p.suffix}")
        print(f"  ! {p.name} 正被其他程序占用(多半是 Excel 或 Adobe 开着)")
        print(f"    改存为 {alt.name}")
        return alt
    except Exception:
        return p


def layout_lines(page, ytol=3.0):
    """pypdf 直接 extract_text 会把 Waybill 的表格打乱(列顺序全错),
       layout 模式对这种 CID 字体又返回空。
       所以自己重建: 用 visitor 拿到每段文字的 x/y 坐标,
       按 y 归成一行、按 x 从左到右排序。"""
    items = []

    def visit(text, cm, tm, font_dict, font_size):
        if text and text.strip():
            items.append((round(tm[5], 1), tm[4], text.strip()))

    try:
        page.extract_text(visitor_text=visit)
    except Exception:
        return []
    items.sort(key=lambda t: (-t[0], t[1]))
    lines, cur, cury = [], [], None
    for y, x, txt in items:
        if cury is None or abs(y - cury) <= ytol:
            cur.append((x, txt))
            cury = y if cury is None else cury
        else:
            lines.append(" ".join(t for _, t in sorted(cur)))
            cur, cury = [(x, txt)], y
    if cur:
        lines.append(" ".join(t for _, t in sorted(cur)))
    return lines


CARGO_ROW = re.compile(
    r"^([A-Z]{4}\d{7})\s+(\d+)\s*[xX]\s*(\d{2}[A-Z]{2})\s+([\d,]+)\s*([A-Z()]+)"
    r"\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)")
SEAL_ROW  = re.compile(r"^SEAL\s+(\S+)")
TOTAL_CNT = re.compile(r"(\d+)\s*CONTAINER\(S\)")


def hbl_items(page):
    """收集每段文字的坐标。tm[4]=x, tm[5]=y(这种 B/L 上 y 向下增大)。
       坐标为 0 的是渲染残留, 丢掉。"""
    out = []

    def v(t, cm, tm, fd, fs):
        if t and t.strip():
            out.append((round(tm[5], 1), round(tm[4], 1), t.strip()))

    try:
        page.extract_text(visitor_text=v)
    except Exception:
        return []
    return sorted(z for z in out if z[0] > 1 and z[1] > 1)


def group_rows(sel, ytol=4):
    """把同一 y 上的碎片拼成一行, 按 x 从左到右。"""
    out, cur, last = [], [], None
    for y, x, t in sorted(sel):
        if last is not None and abs(y - last) > ytol:
            out.append(" ".join(z for _, _, z in sorted(cur, key=lambda a: a[1])))
            cur = []
        cur.append((y, x, t))
        last = y
    if cur:
        out.append(" ".join(z for _, _, z in sorted(cur, key=lambda a: a[1])))
    return out


def parse_hbl(path: Path):
    """从 Transport B/L 抠出 shipper / consignee / notify / marks / description。

       这张单子上 "Shipper" "Consignee or order" 这些标签是背景图不是文字,
       所以认不了标签, 只能靠坐标分区:
         左侧 x<300、港口行以上  -> 按 y 的空档切成三块, 依次是三个当事人
         港口行以下、柜号行以上  -> x<168 是 Marks, 280<=x<560 是 Description

       还有个坑: consignee/notify 写不下时末尾会带一个 *,
       真正的后续内容跑到 Description 下面以 * 开头的那几行去了, 要接回来。"""
    try:
        it = hbl_items(PdfReader(str(path)).pages[0])
    except Exception:
        return {}
    if not it:
        return {}

    left = [z for z in it if z[1] < 300 and z[0] < 460]
    bl, cur, last = [], [], None
    for y, x, t in left:
        if last is not None and y - last > 25:
            bl.append(cur)
            cur = []
        cur.append((y, x, t))
        last = y
    if cur:
        bl.append(cur)

    res = {"_blocks": len(bl)}
    for i, k in enumerate(("shipper", "consignee", "notify")):
        res[k] = "\n".join(t for _, _, t in bl[i]) if i < len(bl) else ""

    cy = [y for y, x, t in it if y > 460 and CONTAINER_PATTERN.fullmatch(t)]
    y1 = min(cy) - 6 if cy else 680
    marks = group_rows([z for z in it if 460 < z[0] < y1 and z[1] < 168])
    desc  = group_rows([z for z in it if 460 < z[0] < y1 and 280 <= z[1] < 560])

    star = next((i for i, l in enumerate(desc) if l.startswith("*")), None)
    cont = []
    if star is not None:
        cont = [l.lstrip("*").strip() for l in desc[star:]]
        desc = desc[:star]
    for k in ("consignee", "notify"):
        if res[k].rstrip().endswith("*"):
            res[k] = res[k].rstrip().rstrip("*") + (("\n" + "\n".join(cont)) if cont else "")

    res["marks"] = "\n".join(marks)
    res["desc"] = "\n".join(desc)
    return res


# Description 区不该出现的东西。出现了说明 x 边界漏进了别的栏位。
STRAY = ("CY-CY", "CY -CY", "SAY TOTAL", "FREIGHT PREPAID", "FREIGHT COLLECT",
         "ON BOARD", "TELEX RELEASE", "AGENT:", "AWELL LOGISTICS")


# 公司名后缀。有些公司名在 B/L 上会折到第二行, 比如
#   NAMSUN(GUANGZHOU) INTERNATIONAL  LOGISTICS
#   CO.,LTD.
# 只取第一行就把 CO.,LTD. 丢了, 所以碰到这种要把下一行接回来。
SUFFIX = re.compile(r"^(CO\s*\.?\s*,?\s*LTD|LTD|INC|LLC|CORP|CORPORATION|COMPANY)\s*\.?\s*,?$")


def split_name(block: str):
    """拆成 (公司名, 余下部分)。公司名 = 第一行; 折行时把续行接回来。"""
    lines = [l.strip() for l in (block or "").split("\n") if l.strip()]
    if not lines:
        return "", ""
    name, used = lines[0], 1
    for nxt in lines[1:]:
        if name.rstrip().endswith(",") or SUFFIX.match(nxt.upper()):
            name = name.rstrip().rstrip(",") + " " + nxt
            used += 1
        else:
            break
    return re.sub(r"\s{2,}", " ", name).strip(), "\n".join(lines[used:])


def first_name(block: str) -> str:
    return split_name(block)[0]


def norm_name(n: str) -> str:
    """比较公司名用。去掉标点和多余空格, 免得 CO.,LTD. 和 CO., LTD 被当成两家。"""
    return re.sub(r"[^A-Z0-9]", "", (n or "").upper())


def check_hbl(party, nblocks):
    """B/L 解析的自检。
       切块靠行距、Marks/Description 靠 x 坐标画框, 都有可能失手,
       所以把能自动发现的异常都列出来, 让人去看那几票, 而不是全部人工核对。"""
    bad = []
    # 左侧除了三个当事人, 下面还有船名/港口两行, 所以块数通常是 4~5。
    # 取前三块当当事人, 只有不足 3 块才是真出问题。
    if nblocks < 3:
        bad.append(f"左侧只切出 {nblocks} 块, 三个当事人没分全")
    for k, name in (("shipper", "Shipper"), ("consignee", "Consignee"),
                    ("notify", "Notify"), ("marks", "Marks"),
                    ("desc", "Description")):
        if not (party.get(k) or "").strip():
            bad.append(f"{name} 是空的")
    desc = (party.get("desc") or "").upper()
    hit = [w for w in STRAY if w in desc]
    if hit:
        bad.append("Description 里混进了 " + "/".join(hit))
    if CONTAINER_PATTERN.search(desc.replace("\n", " ")):
        bad.append("Description 里混进了柜号")
    # 当事人栏里混进船名/港口 -> 说明块边界串了
    for k, name in (("shipper", "Shipper"), ("consignee", "Consignee"), ("notify", "Notify")):
        up = (party.get(k) or "").upper()
        if "CMA CGM" in up or "PORT OF" in up:
            bad.append(f"{name} 里混进了船名或港口")
    for k, name in (("shipper", "Shipper"), ("consignee", "Consignee"), ("notify", "Notify")):
        if (party.get(k) or "").count("\n") >= 9:
            bad.append(f"{name} 有 {party[k].count(chr(10))+1} 行, 可能把下一块也吞进来了")
    return bad


def parse_waybill(path: Path):
    """从 Waybill 第 1 页抠货物信息。
       版面长这样:
           CMAU7833997   1 x 40HC   1018 CARTONS   16619.000   3700   68.000
           SEAL R9073324
           Weight in Kgs Total: 1 CONTAINER(S)
       返回 {柜数, 柜号, 封条, 柜型, 件数, 重量, 材积}, 失败返回 {}"""
    try:
        lines = layout_lines(PdfReader(str(path)).pages[0])
    except Exception:
        return {}

    cons, seals, declared = [], [], ""
    for raw in lines:
        for l in raw.split("\n"):
            l = l.strip()
            m = CARGO_ROW.match(l)
            if m and not any(c[0] == m.group(1) for c in cons):
                cons.append(m.groups())
            m2 = SEAL_ROW.match(l)
            if m2 and m2.group(1) not in seals:
                seals.append(m2.group(1))
            m3 = TOTAL_CNT.search(l)
            if m3 and not declared:
                declared = m3.group(1)

    if not cons:
        return {}
    j = lambda i: ",".join(c[i] for c in cons)
    return {
        "柜数": declared or str(len(cons)),
        "柜号": j(0),
        "封条": ",".join(seals) if seals else "",
        "柜型": j(2),
        "件数": j(3),            # 单位(CARTONS)已经分开, 不带进来
        "重量": j(5),
        "材积": j(7),
    }


def manifest_rows(path: Path):
    """读舱单, 返回 {MBL: {...}}。不认列名(船公司表头会变), 靠内容判断每一列是什么。"""
    rows = []
    suf = path.suffix.lower()
    try:
        if suf == ".xls":
            import xlrd
            bk = xlrd.open_workbook(str(path))
            for sh in bk.sheets():
                for r in range(sh.nrows):
                    rows.append([str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)])
        elif suf in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            for ws in load_workbook(str(path), read_only=True, data_only=True).worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append([("" if v is None else str(v)).strip() for v in row])
        else:
            import csv as _csv
            with open(path, encoding="utf-8-sig", errors="ignore") as fh:
                rows = [r for r in _csv.reader(fh)]
    except ImportError:
        print(f"  ! 读不了 {path.name}: 缺少 xlrd, 请执行  pip install xlrd")
        return {}
    except Exception as e:
        print(f"  ! 读不了 {path.name}: {e}")
        return {}

    SIZE = re.compile(r"(\d+)\s*[xX]\s*(\d{2}[A-Z]{2})")
    out = {}
    for r in rows:
        mbl = next((v.upper() for v in r
                    if KEY_PATTERN.fullmatch(v.upper()) and not CONTAINER_PATTERN.fullmatch(v.upper())), "")
        if not mbl:
            continue
        hbl = next((v.upper() for v in r if HBL_PATTERN.search(v.upper())), "")
        sz  = next((SIZE.search(v) for v in r if SIZE.search(v)), None)
        pods = [v for v in r if v.isupper() and v.replace(" ", "").isalpha() and len(v) > 3]
        out.setdefault(mbl, []).append({
            "HBL舱单": hbl,
            "柜数舱单": int(sz.group(1)) if sz else "",
            "柜型舱单": sz.group(2) if sz else "",
            "始发港": pods[0] if len(pods) > 0 else "",
            "目的港": pods[1] if len(pods) > 1 else "",
            "备注": next((v for v in r if any(ch > "\u4e00" for ch in v)), ""),
            "文件包": next((v for v in r if v.upper().startswith("PART")), ""),
        })
    return out


def pick_manifest(records, hbl):
    """一个 MBL 在舱单里可能占好几行(分单给不同客户), 按 HBL 挑出对应那行。
       舱单上的号带 AIC 前缀(AIC2610A334703A), 文件名里是裸号(2610A334703A),
       所以用包含判断而不是相等。"""
    if not records:
        return {}
    if len(records) == 1:
        return records[0]
    if hbl:
        for r in records:
            if hbl in (r.get("HBL舱单") or ""):
                return r
    return records[0]


def load_manifest(path: Path):
    """读船公司给的舱单(xls/xlsx/csv), 把里面所有 MBL 号抠出来。
       不认列名, 直接全表扫描找符合提单号格式的格子 —— 船公司的表头经常变,
       认列名很容易哪天就挂了, 扫格子反而稳。"""
    cells = []
    suf = path.suffix.lower()
    try:
        if suf == ".xls":
            import xlrd
            bk = xlrd.open_workbook(str(path))
            for sh in bk.sheets():
                for r in range(sh.nrows):
                    for c in range(sh.ncols):
                        cells.append(str(sh.cell_value(r, c)))
        elif suf in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells.extend(str(v) for v in row if v is not None)
        else:                                   # csv / txt
            with open(path, encoding="utf-8-sig", errors="ignore") as fh:
                for line in fh:
                    cells.extend(line.replace(",", " ").split())
    except ImportError:
        print(f"  ! 读不了 {path.name}: 缺少 xlrd, 请执行  pip install xlrd")
        return set()
    except Exception as e:
        print(f"  ! 读不了 {path.name}: {e}")
        return set()

    keys = set()
    for v in cells:
        for m in KEY_PATTERN.finditer(v.upper()):
            k = m.group(0)
            if not CONTAINER_PATTERN.fullmatch(k):
                keys.add(k)
    return keys


SCAC = "CMDU"            # MBL 前缀, 换船公司时改这里

S1_FIELDS = [
    ("MBL",      "MBL"),
    ("HBL",      "HBL"),
    ("柜数",     "柜数"),
    ("Shipper Name",       "ShipperName"),
    ("Shipper",            "Shipper"),
    ("Consignee Name",     "ConsigneeName"),
    ("Consignee or order", "Consignee"),
    ("Notify Name",        "NotifyName"),
    ("Notify address",     "Notify"),
    ("柜号",     "柜号"),
    ("封条",     "封条"),
    ("柜型",     "柜型"),
    ("件数",     "件数"),
    ("重量",     "重量"),
    ("材积",     "材积"),
    ("Marks and Numbers",   "Marks"),
    ("Description of Goods","Description"),
    ("运输方式", "运输方式"),
    ("起始港",   "起始港"),
    ("目的地",   "目的地"),
    ("ETA",      "ETA"),
    ("状态",     "状态"),
    ("第几张纸", "第几张纸"),
    ("批次",     "批次"),
    ("文件包",   "文件包"),
    ("备注",     "备注"),
    ("BL文件",   "BL文件"),
    ("Waybill文件", "Waybill文件"),
]


def strip_unit(v):
    """1041CARTONS -> 1041 ; 13457.000 KGS -> 13457.000 ; 68.000CBM -> 68.000"""
    if not v:
        return ""
    parts = [re.sub(r"[A-Za-z()\s]+$", "", x).strip() for x in str(v).split(",")]
    return ",".join(p for p in parts)


def write_s1_report(out_dir, rows, problems, man_lines):
    """第一叠报告。xlsx 带高亮, 没装 openpyxl 就退回 csv。
       Ramp 整行红字, 柜型不是 40HC 的那一格标橙。"""
    headers = [h for h, _ in S1_FIELDS]
    keys    = [k for _, k in S1_FIELDS]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter as L
    except ImportError:
        print("  ! 没装 openpyxl, 报告存成 csv(无高亮)。想要高亮:  pip install openpyxl")
        path = writable_path(out_dir / "_报告.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(headers)
            for r in rows:
                w.writerow([r.get(k, "") for k in keys])
            if man_lines:
                fh.write("\n舱单核对\n")
                for a, b in man_lines:
                    fh.write(f'"{a}","{b}"\n')
            fh.write("\n问题清单\n文件,原因\n")
            for n, why in problems:
                fh.write(f'"{n}","{why}"\n')
        return path

    ORANGE = PatternFill("solid", fgColor="FCE4D6")
    PURPLE = PatternFill("solid", fgColor="E4DFEC")   # Notify 里没有邮箱
    YELLOW = PatternFill("solid", fgColor="FFF2CC")   # Notify 公司名跟 Consignee 不同
    HEAD   = PatternFill("solid", fgColor="DDEBF7")
    RED    = Font(color="C00000", bold=True)   # Ramp
    VIOLET = Font(color="7E57C2", bold=True)   # 跟上一行共用一张纸的分单
    BLUEFILL = PatternFill("solid", fgColor="DDEBF7")   # 起始港跟大多数不一样
    LILAC    = PatternFill("solid", fgColor="EDE7F6")   # 这一行有问题, 整行淡紫

    wb = Workbook(); ws = wb.active; ws.title = "明细"
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).font = Font(bold=True); ws.cell(1, c).fill = HEAD

    from openpyxl.styles import Alignment
    size_col   = keys.index("柜型") + 1
    notify_col = keys.index("Notify") + 1
    nname_col  = keys.index("NotifyName") + 1
    wrap_cols  = [keys.index(k) + 1 for k in
                  ("Shipper", "Consignee", "Notify", "Marks", "Description")]
    for r in rows:
        ws.append([("" if r.get(k) in (None,) else r.get(k)) for k in keys])
        i = ws.max_row
        # Ramp 整行红字(字的颜色), 起始港异常整行蓝底(背景色)。
        # 一个管字一个管底, 所以两个都中也不冲突, 会同时显示。
        if str(r.get("运输方式", "")).lower() == "ramp":
            for c in range(1, len(headers) + 1):
                ws.cell(i, c).font = RED
        if r.get("DupRow"):     # 不占纸的分单行 -> 淡紫字, 压过 Ramp 的红字
            for c in range(1, len(headers) + 1):
                ws.cell(i, c).font = VIOLET
        if r.get("PolOdd"):
            for c in range(1, len(headers) + 1):
                ws.cell(i, c).fill = BLUEFILL
        if r.get("Issue"):          # 有问题的整行淡紫, 压过起始港的蓝底
            for c in range(1, len(headers) + 1):
                ws.cell(i, c).fill = LILAC
        for c in wrap_cols:                       # 多行内容要自动换行才看得见
            ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")
        # 跟 Consignee 不一样的那一格标黄: 名字归名字, 地址归地址
        if r.get("NotifyDiff"):
            ws.cell(i, nname_col).fill = YELLOW
        if r.get("NotifyAddrDiff"):
            ws.cell(i, notify_col).fill = YELLOW
        elif not r.get("NotifySameAs") and "@" not in str(r.get("Notify", "")):
            ws.cell(i, notify_col).fill = PURPLE
        sz = str(r.get("柜型", ""))
        if sz and any(x.strip().upper() != "40HC" for x in sz.split(",")):
            ws.cell(i, size_col).fill = ORANGE

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{L(len(headers))}{ws.max_row}"
    for c, (h, k) in enumerate(S1_FIELDS, 1):
        if k in ("Shipper", "Consignee", "Notify", "Marks", "Description"):
            ws.column_dimensions[L(c)].width = 42
            continue
        longest = max([len(str(h))] + [len(str(x.get(k, ""))) for x in rows] or [0])
        ws.column_dimensions[L(c)].width = min(max(longest + 2, 8), 40)

    w2 = wb.create_sheet("核对与异常")
    w2.append(["类型", "内容"])
    for c in (1, 2):
        w2.cell(1, c).font = Font(bold=True); w2.cell(1, c).fill = HEAD
    for a, b in man_lines:
        w2.append([a, b])
    for n, why in problems:
        w2.append([n, why])
    w2.column_dimensions["A"].width = 34
    w2.column_dimensions["B"].width = 60

    ramp = sum(1 for r in rows if str(r.get("运输方式", "")).lower() == "ramp")
    odd  = sum(1 for r in rows if r.get("柜型") and
               any(x.strip().upper() != "40HC" for x in str(r["柜型"]).split(",")))
    ndiff  = sum(1 for r in rows if r.get("NotifyDiff") or r.get("NotifyAddrDiff"))
    ndup = sum(1 for r in rows if r.get("DupRow"))
    if ndup:
        print(f"  同 MBL 的分单 {ndup} 行 -> 不单独占纸, 第几张纸留空 + 淡紫字")
    nissue = sum(1 for r in rows if r.get("Issue"))
    if nissue:
        print(f"  有问题需要看的 {nissue} 票 -> 整行淡紫, 详见状态列")
    npol   = sum(1 for r in rows if r.get("PolOdd"))
    if npol:
        print(f"  起始港跟大多数不同 {npol} 票 -> 整行蓝底")
    nomail = sum(1 for r in rows if not r.get("NotifyDiff") and not r.get("NotifySameAs")
                 and "@" not in str(r.get("Notify", "")))
    if ndiff:
        print(f"  Notify 跟 Consignee 不一致 {ndiff} 票 -> 不一致的那格标黄")
    print(f"  Ramp(转铁路) {ramp} 票 -> 整行红字")
    if nomail:
        print(f"  Notify 里没有邮箱 {nomail} 票 -> 该格标紫")
    print(f"  非 40HC 柜型 {odd} 票 -> 柜型格标橙")

    path = writable_path(out_dir / "_报告.xlsx")
    wb.save(path)
    return path


def is_output_file(f: Path) -> bool:
    return f.is_file() and (f.name.startswith(OUT_PREFIX) or f.name in EXTRA_OUT)


def archive_old_output(out_dir: Path, archive_root: Path):
    """把上一批留下的成品挪走, 跟那一批的原始文件归到同一个文件夹"""
    olds = [f for f in out_dir.iterdir() if is_output_file(f)]
    marker = out_dir / MARKER

    if not olds:
        if marker.exists():
            marker.unlink()
        return

    if archive_root is None:                      # 没配归档就只能删
        for f in olds:
            try:
                f.unlink()
            except Exception:
                pass
        print(f"清掉上一批的 {len(olds)} 个成品文件")
        return

    name = ""
    if marker.exists():
        try:
            name = marker.read_text(encoding="utf-8").strip()
        except Exception:
            name = ""
    if not name:                                  # 记号丢了就用文件修改时间兜底
        name = datetime.fromtimestamp(olds[0].stat().st_mtime).strftime("%Y-%m-%d_%H%M")

    dest = archive_root / name
    dest.mkdir(parents=True, exist_ok=True)
    moved = 0
    for f in olds:
        try:
            t = dest / f.name
            i = 1
            while t.exists():
                t = dest / f"{f.stem}({i}){f.suffix}"
                i += 1
            shutil.move(str(f), str(t))
            moved += 1
        except Exception as e:
            print(f"  ! 挪不动 {f.name}: {e}")
    print(f"上一批的 {moved} 个成品已归档 -> {name}")

    if marker.exists():
        marker.unlink()


def unzip_all(folder: Path):
    zips = list(folder.glob("*.zip"))
    if not zips:
        return
    n = 0
    for z in zips:
        try:
            with zipfile.ZipFile(z) as zf:
                for info in zf.infolist():
                    if info.is_dir():
                        continue
                    name = Path(info.filename).name
                    if not name.lower().endswith(".pdf"):
                        continue
                    target = folder / name
                    if target.exists():
                        continue
                    with zf.open(info) as s, open(target, "wb") as d:
                        d.write(s.read())
                    n += 1
        except Exception as e:
            print(f"  ! 解压失败 {z.name}: {e}")
    print(f"解压 {len(zips)} 个 zip,释放 {n} 个新 PDF")


def read_pdf(path: Path):
    try:
        r = PdfReader(str(path))
        return len(r.pages), (r.pages[0].extract_text() or "").upper()
    except Exception as e:
        print(f"  ! 读不了 {path.name}: {e}")
        return 0, ""


def kind_by_name(path: Path):
    """返回 (MBL号, 类型, HBL号)
       CHN3962684 2609A336281 -> ("CHN3962684", "bl", "2609A336281")
       CHN3962684             -> ("CHN3962684", "wb", "")
    """
    stem = path.stem.upper().strip()
    m = KEY_PATTERN.search(stem)
    if not m:
        return None, None, ""
    rest = (stem[:m.start()] + stem[m.end():]).strip(" _-")
    hbl = re.split(r"[\s_\-]+", rest)[0] if rest else ""
    return m.group(0), ("bl" if rest else "wb"), hbl


def kind_by_text(txt: str):
    if BL_HINT in txt:
        return "bl"
    if WB_HINT in txt:
        return "wb"
    return None


def archive_sources(src: Path, archive_root: Path) -> str:
    """把处理过的 zip 和 pdf 挪到归档文件夹, 返回文件夹名"""
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    dest = archive_root / stamp
    dest.mkdir(parents=True, exist_ok=True)
    moved = 0
    picks = [f for f in src.iterdir()
             if f.is_file() and f.suffix.lower() in (".zip", ".pdf", ".xls", ".xlsx", ".xlsm")]
    for f in picks:
        try:
            t = dest / f.name
            i = 1
            while t.exists():
                t = dest / f"{f.stem}({i}){f.suffix}"
                i += 1
            shutil.move(str(f), str(t))
            moved += 1
        except Exception as e:
            print(f"  ! 归档失败 {f.name}: {e}")
    print(f"已归档 {moved} 个原始文件 -> {stamp}")
    return stamp


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", help="输入文件夹 1_DROP_ZIPS_HERE")
    ap.add_argument("--out", default=None, help="输出文件夹 2_READY_TO_PRINT")
    ap.add_argument("--archive", default=None, help="归档文件夹 3_ARCHIVE")
    ap.add_argument("--chunk", type=int, default=0, help="每个输出文件放几票, 0=不拆")
    ap.add_argument("--manifest", default=None,
                    help="船公司舱单 xls/xlsx/csv, 跑完自动核对缺哪几票")
    ap.add_argument("--dry-run", action="store_true", help="只检查, 什么都不动")
    args = ap.parse_args()

    src = Path(args.folder)
    if not src.is_dir():
        sys.exit(f"找不到输入文件夹: {src}")
    out_dir = Path(args.out) if args.out else src / "_PRINT"
    out_dir.mkdir(parents=True, exist_ok=True)
    arc = Path(args.archive) if args.archive else None
    if arc:
        arc.mkdir(parents=True, exist_ok=True)

    if not args.dry_run:
        archive_old_output(out_dir, arc)

    unzip_all(src)

    pdfs = sorted(p for p in src.glob("*.pdf") if p.is_file())
    print(f"共 {len(pdfs)} 个 PDF\n")
    if not pdfs:
        print("输入文件夹里没有 PDF。请先把 zip 放进 1_DROP_ZIPS_HERE。")
        return

    # 一票 = 一个 HBL。同一个 MBL 下可以挂多个 HBL(分单给不同客户),
    # 所以 B/L 按 (MBL, HBL) 存, Waybill 按 MBL 存 —— 船公司的 Waybill 是
    # 按主提单发的, 一份会被同 MBL 下的几票共用。
    bl_items, wb_items, problems = {}, {}, []
    for p in pdfs:
        key, k_name, hbl = kind_by_name(p)
        npages, txt = read_pdf(p)
        k_text = kind_by_text(txt)
        if not key:
            problems.append((p.name, "文件名里找不到提单号"))
            continue
        kind = k_name or k_text
        if kind is None:
            problems.append((p.name, "认不出是 B/L 还是 Waybill"))
            continue
        if k_text and k_name and k_text != k_name:
            problems.append((p.name, f"文件名像 {k_name} 但内容像 {k_text},已按文件名处理"))
        if kind == "bl":
            if not hbl:
                problems.append((p.name, f"{key}: B/L 文件名里没找到 HBL 号,排到最后"))
            if (key, hbl) in bl_items:
                problems.append((p.name, f"{key} {hbl}: 同一个 HBL 有两份 B/L, 用了先读到的那份"))
                continue
            bl_items[(key, hbl)] = (p, npages)
        else:
            if key in wb_items:
                problems.append((p.name, f"{key}: 同一个 MBL 有两份 Waybill, 用了先读到的那份"))
                continue
            wb_items[key] = (p, npages)

    # 配不齐的(只有 B/L 没 Waybill, 或反过来)也照常进队列, 缺的那面补白页。
    # 这样整叠纸的顺序不会因为少一票而乱, 报告里也看得到它。
    matched, singles = [], []
    used_wb = set()
    for (key, hbl), bl in bl_items.items():
        wb = wb_items.get(key)
        if wb:
            used_wb.add(key)
        else:
            singles.append((key, bl[0]))
            problems.append((bl[0].name, f"{key}: 缺 Waybill(反面留白)"))
        matched.append((key, hbl, bl, wb))
    for key, wb in wb_items.items():          # 只有 Waybill 没 B/L 的
        if key not in used_wb:
            singles.append((key, wb[0]))
            problems.append((wb[0].name, f"{key}: 缺 Transport B/L(正面留白)"))
            matched.append((key, "", None, wb))

    # 按 HBL 排序; 没有 HBL 的排最后, 再按 MBL 排
    matched.sort(key=lambda x: (x[1] == "", x[1], x[0]))
    singles.sort(key=lambda x: x[0])

    # 同一个 MBL 下有几票 —— 大于 1 就是分单, 这几票共用一份 Waybill
    mbl_rows = Counter(m for m, _, _, _ in matched)

    print(f"共 {len(matched)} 票 | 其中配不齐的 {len(singles)} 票(缺的那面补白页) | 问题 {len(problems)}\n")

    # ---- 跟船公司舱单核对 ----
    man_lines, man_info = [], {}
    mp = Path(args.manifest) if args.manifest else None
    if mp is None:                                   # 没指定就在输入文件夹里自己找
        found = sorted(f for f in src.iterdir()
                       if f.is_file() and f.suffix.lower() in (".xls", ".xlsx", ".xlsm"))
        if found:
            mp = found[0]
            print(f"发现舱单: {mp.name}\n")
    if mp:
        if not mp.exists():
            print(f"!! 找不到舱单: {mp}\n")
        else:
            man = manifest_rows(mp)
            man_info.update(man)
            want = set(man) or load_manifest(mp)
            have = {k for k, _, _, _ in matched} | {k for k, _ in singles}
            lack = sorted(want - have)
            more = sorted(have - want)
            print(f"--- 舱单核对 ({mp.name}) ---")
            man_lines_n = sum(len(v) for v in man.values())
            if man_lines_n != len(want):
                print(f"  舱单共      {len(want)} 个主提单 / {man_lines_n} 张分单")
            else:
                print(f"  舱单共      {len(want)} 票")
            print(f"  手上有      {len(have)} 票")
            print(f"  舱单有你没有 {len(lack)} 票")
            print(f"  你有舱单没有 {len(more)} 票")
            if lack:
                print("  缺这些:")
                for k in lack:
                    print(f"     {k}")
            if more:
                print("  多这些:")
                for k in more:
                    print(f"     {k}")
            print()
            man_lines.append(("舱单核对", mp.name))
            man_lines.append(("舱单票数", len(want)))
            if man_lines_n != len(want):
                man_lines.append(("舱单分单行数", man_lines_n))
            man_lines.append(("手上票数", len(have)))
            for k in lack:
                man_lines.append(("舱单有你没有", k))
            for k in more:
                man_lines.append(("你有舱单没有", k))

    main_pod = main_pol = ""
    if man_info:
        recs = [r for rs in man_info.values() for r in rs]
        pods = [r.get("目的港", "") for r in recs if r.get("目的港")]
        if pods:
            main_pod = max(set(pods), key=pods.count)
        pols = [r.get("始发港", "") for r in recs if r.get("始发港")]
        if pols:
            main_pol = max(set(pols), key=pols.count)

    rows = []
    batches = ([matched] if args.chunk <= 0
               else [matched[i:i + args.chunk] for i in range(0, len(matched), args.chunk)])

    printed = set()     # 已经出过纸的 MBL
    for bi, batch in enumerate(batches, 1):
        writer = None if args.dry_run else PdfWriter()
        sheet = 0
        for key, hbl, blp, wbp in batch:
            bl, nbl = blp if blp else (None, 0)
            wb, nwb = wbp if wbp else (None, 0)
            take = [i for i in WAYBILL_PAGES if i <= nwb]
            total = max(nbl, 1) + max(len(take), 1)      # 缺哪面就用一张白页顶上
            pad = 1 if total % 2 else 0
            # 同一个 MBL 下的几个分单是同一票货(同柜同封条), 只出一张纸。
            # 后面的分单照样在明细里占一行, 只是「第几张纸」留空 —— 两叠纸都按
            # 这个规则走, 张数才对得上。
            dup = key in printed
            if dup:
                start_sheet = ""
            else:
                printed.add(key)
                if writer is not None:
                    if bl:
                        for pg in PdfReader(str(bl)).pages:
                            writer.add_page(pg)
                    else:
                        writer.add_blank_page(width=612, height=792)
                    if wb and take:
                        wbr = PdfReader(str(wb))
                        for i in take:
                            writer.add_page(wbr.pages[i - 1])
                    else:
                        writer.add_blank_page(width=612, height=792)
                    if pad:
                        last = writer.pages[-1]
                        writer.add_blank_page(width=last.mediabox.width,
                                              height=last.mediabox.height)
                start_sheet = sheet + 1
                sheet += (total + pad) // 2
            # 这一行要说的话。hard=True 是真的数据问题 —— 明细里整行标淡紫,
            # 提醒人去看。hard=False 是提示: 起始港有蓝底、Notify 有黄格,
            # 已经各自标过了, 再把整行也染上反而看不出哪几行真的要处理。
            msgs = []

            def say(m, hard=True):
                msgs.append((m, hard))

            if not bl:
                say("没有 Transport B/L, 正面是白页")
            elif nbl != 1:
                say(f"B/L 有 {nbl} 页, 正反面可能错位")
            if not wb:
                say("没有 Waybill, 反面是白页")
            elif nwb < max(WAYBILL_PAGES):
                say(f"Waybill 只有 {nwb} 页")
            if mbl_rows[key] > 1:
                say(f"同一个 MBL 下有 {mbl_rows[key]} 个分单, 共用一张纸"
                    + ("(本行不单独占纸)" if dup else ""))

            info = pick_manifest(man_info.get(key, []), hbl)
            # HBL 用舱单上的托单编号(带 AIC 前缀), 舱单里没有就退回文件名里的号
            hbl_show = info.get("HBL舱单", "") or hbl
            if info.get("HBL舱单"):
                if hbl and hbl not in info["HBL舱单"]:
                    say(f"HBL 对不上: 文件名 {hbl}, 舱单 {info['HBL舱单']}")
            elif info:
                say("舱单上没有托单编号, HBL 用的是文件名里的号")
            elif man_info:
                say("这票不在舱单里")

            party = parse_hbl(bl) if bl else {}
            sh_name, sh_rest = split_name(party.get("shipper", ""))
            cn_name, cn_rest = split_name(party.get("consignee", ""))
            nt_name, nt_rest = split_name(party.get("notify", ""))
            # Notify 正常情况跟 Consignee 一模一样。名字和地址分开比,
            # 哪一半不一样就标哪一格 —— 地址少一行多一行也算不一样。
            nt_same_as = "SAMEASCONSIGNEE" in norm_name(party.get("notify", ""))
            if nt_same_as:
                nt_name_show, nt_show = "", party.get("notify", "")
                nt_diff = nt_addr_diff = False
            else:
                nt_name_show, nt_show = nt_name, nt_rest
                nt_diff = norm_name(nt_name) != norm_name(cn_name)
                nt_addr_diff = norm_name(nt_rest) != norm_name(cn_rest)
            if nt_diff:
                say(f"Notify 公司名跟 Consignee 不同: {nt_name}", False)
            if nt_addr_diff:
                say("Notify 地址跟 Consignee 不同", False)
            if bl and not party:
                say("B/L 读不出来")
            elif party:
                for w in check_hbl(party, party.get("_blocks", 0)):
                    say(w)

            cargo = parse_waybill(wb) if wb else {}
            if wb and not cargo:
                say("Waybill 货物信息读不出来")
            elif cargo:     # 压根没有 Waybill 时上面已经说过了, 别再报一次柜数对不上
                # 跟舱单的箱型箱量对一下, 多柜最容易在这里出岔子
                want_n = info.get("柜数舱单", "")
                got_n = cargo.get("柜号", "").count(",") + 1 if cargo.get("柜号") else 0
                if want_n and int(want_n) != got_n:
                    say(f"柜数对不上: 舱单 {want_n} 个, Waybill 只读出 {got_n} 个")
                want_t = (info.get("柜型舱单", "") or "").upper()
                got_t = {x.strip().upper() for x in cargo.get("柜型", "").split(",") if x.strip()}
                if want_t and got_t and want_t not in got_t:
                    say(f"柜型对不上: 舱单 {want_t}, Waybill {'/'.join(sorted(got_t))}")

            dest = info.get("目的港", "")
            pol = info.get("始发港", "")
            # 绝大多数是同一个起始港, 冒出别的就单独拎出来
            pol_odd = bool(pol) and bool(main_pol) and pol != main_pol
            if pol_odd:
                say(f"起始港不是 {main_pol}, 是 {pol}", False)
            mode = ""
            if dest:
                mode = "Ramp" if dest != main_pod else "Port"

            note = " / ".join(m for m, _ in msgs)
            issue = any(hard for _, hard in msgs)

            rows.append({
                "MBL": (SCAC + key) if key else "",
                "HBL": hbl_show,
                "ShipperName": sh_name,
                "Shipper": sh_rest,
                "ConsigneeName": cn_name,
                "Consignee": cn_rest,
                "NotifyName": nt_name_show,
                "Notify": nt_show,
                "NotifyDiff": nt_diff,
                "NotifyAddrDiff": nt_addr_diff,
                "NotifySameAs": nt_same_as,
                "Marks": party.get("marks", ""),
                "Description": party.get("desc", ""),
                "柜数": cargo.get("柜数", "") or info.get("柜数舱单", "") or "",
                "柜号": cargo.get("柜号", ""),
                "封条": cargo.get("封条", ""),
                "柜型": cargo.get("柜型", "") or info.get("柜型舱单", "") or "",
                "件数": cargo.get("件数", ""),
                "重量": cargo.get("重量", ""),
                "材积": cargo.get("材积", ""),
                "运输方式": mode, "起始港": pol, "目的地": dest, "ETA": "",
                "PolOdd": pol_odd,
                "状态": note or "正常",
                "Issue": issue,
                "DupRow": dup,
                "第几张纸": start_sheet, "批次": bi,
                "文件包": info.get("文件包", ""), "备注": info.get("备注", ""),
                "BL文件": bl.name if bl else "", "Waybill文件": wb.name if wb else "",
            })

        if writer is not None:
            name = "PRINT.pdf" if len(batches) == 1 else f"PRINT_{bi:02d}.pdf"
            with open(writable_path(out_dir / name), "wb") as fh:
                writer.write(fh)
            print(f"  {name}: {len(writer.pages)} 页 = {len(writer.pages)//2} 张纸")

    report_path = write_s1_report(out_dir, rows, problems, man_lines)

    print(f"\n报告: {report_path}")

    if args.dry_run:
        print("\n(检查模式: 没有生成 PDF, 什么文件都没动)")
    elif arc:
        stamp = archive_sources(src, arc)
        # 留个记号, 下次跑的时候把这批成品挪到同一个归档文件夹
        (out_dir / MARKER).write_text(stamp, encoding="utf-8")


if __name__ == "__main__":
    main()
